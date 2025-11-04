#!/usr/bin/env python3
"""
Resume Contact Information Extractor
Automates extraction of contact details from resumes and populates Excel spreadsheet.
"""

import json
import logging
import sys
from pathlib import Path
from typing import Dict, Optional
import re
import pandas as pd
from openpyxl import load_workbook
import docx
from datetime import datetime

# Configure logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[
        logging.FileHandler('resume_extractor.log'),
        logging.StreamHandler(sys.stdout)
    ]
)

class ResumeExtractor:
    def __init__(self, config_path: str = "config.json"):
        self.config = self._load_config(config_path)
        self.output_file = Path(self.config["output_excel_path"])
        self.resume_folder = Path(self.config["resume_folder"])

    def _load_config(self, config_path: str) -> dict:
        """Load configuration from JSON file."""
        try:
            with open(config_path, 'r') as f:
                return json.load(f)
        except FileNotFoundError:
            # Create default config if not exists
            default_config = {
                "resume_folder": "resumes",
                "output_excel_path": "contacts.xlsx",
                "required_fields": ["name", "email", "phone"]
            }
            with open(config_path, 'w') as f:
                json.dump(default_config, f, indent=4)
            return default_config

    def extract_contact_info(self, doc_text: str) -> Dict[str, str]:
        """Extract contact information from resume text using regex patterns."""
        patterns = {
            'name': r'^([A-Z][a-z]+(?:\s[A-Z][a-z]+)+)',
            'email': r'[\w\.-]+@[\w\.-]+\.\w+',
            'phone': r'(\+\d{1,2}\s)?\(?\d{3}\)?[\s.-]\d{3}[\s.-]\d{4}'
        }
        
        contact_info = {}
        for field, pattern in patterns.items():
            match = re.search(pattern, doc_text)
            contact_info[field] = match.group(0) if match else ''
            
        return contact_info

    def process_resume(self, resume_path: Path) -> Optional[Dict[str, str]]:
        """Process a single resume document and extract contact information."""
        try:
            doc = docx.Document(resume_path)
            text = '\n'.join([paragraph.text for paragraph in doc.paragraphs])
            contact_info = self.extract_contact_info(text)
            contact_info['filename'] = resume_path.name
            contact_info['processed_date'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
            return contact_info
        except Exception as e:
            logging.error(f"Error processing {resume_path}: {str(e)}")
            return None

    def update_excel(self, contact_info: Dict[str, str]) -> None:
        """Update Excel spreadsheet with extracted contact information."""
        try:
            if not self.output_file.exists():
                df = pd.DataFrame(columns=list(contact_info.keys()))
                df.to_excel(self.output_file, index=False)

            df = pd.read_excel(self.output_file)

            # ✅ Fix: use concat instead of append
            new_row = pd.DataFrame([contact_info])
            df = pd.concat([df, new_row], ignore_index=True)

            df.to_excel(self.output_file, index=False)
            logging.info(f"Updated contact information for {contact_info['name']}")
        except Exception as e:
            logging.error(f"Error updating Excel file: {str(e)}")

    def run(self) -> None:
        """Main execution method to process all resumes in the folder."""
        if not self.resume_folder.exists():
            self.resume_folder.mkdir(parents=True)
            logging.info(f"Created resume folder: {self.resume_folder}")
            return

        for resume_path in self.resume_folder.glob("*.docx"):
            logging.info(f"Processing resume: {resume_path.name}")
            contact_info = self.process_resume(resume_path)
            if contact_info:
                self.update_excel(contact_info)

def main():
    try:
        extractor = ResumeExtractor()
        extractor.run()
    except Exception as e:
        logging.error(f"Application error: {str(e)}")
        sys.exit(1)

if __name__ == "__main__":
    main()